"""
Calcula las metricas del objetivo especifico 2 a partir de la anotacion manual.

Lee los tres libros completados en anotacion/, mide la concordancia entre
anotadores, consolida un estandar de oro por mayoria y evalua contra el el
desempeño del sistema.

Uso:
    ./venv/Scripts/python.exe evaluar_anotacion.py
    ./venv/Scripts/python.exe evaluar_anotacion.py --salida resultados.md
"""
import argparse
import itertools
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CARPETA = Path("anotacion")
ANOTADORES = ["Daniela_Trujillo", "Miguel_Reina", "Gabriel_Zambrano"]
PRIMERA_FILA = 7          # las filas 1-5 son encabezado y la 6 es el ejemplo
VALIDAS = {"SI", "NO", "DUDOSO"}


# ---------------------------------------------------------------- lectura
def leer_hoja(ruta: Path, hoja: str, col_id: int = 1, col_valor: int = 3) -> dict:
    """Devuelve {id: etiqueta} de una hoja de anotacion."""
    wb = load_workbook(ruta, data_only=True)
    if hoja not in wb.sheetnames:
        return {}
    h = wb[hoja]
    out = {}
    for fila in range(PRIMERA_FILA, h.max_row + 1):
        ident = h.cell(row=fila, column=col_id).value
        valor = h.cell(row=fila, column=col_valor).value
        if ident is None:
            continue
        etiqueta = str(valor).strip().upper() if valor is not None else ""
        out[str(ident).strip()] = etiqueta if etiqueta in VALIDAS else ""
    return out


def cargar(hoja: str) -> pd.DataFrame:
    """Matriz items x anotadores."""
    columnas = {}
    for nombre in ANOTADORES:
        ruta = CARPETA / f"Anotacion_{nombre}.xlsx"
        if not ruta.exists():
            print(f"  aviso: falta {ruta.name}, se omite ese anotador.")
            continue
        columnas[nombre] = leer_hoja(ruta, hoja)
    if not columnas:
        return pd.DataFrame()
    return pd.DataFrame(columnas).sort_index()


# ---------------------------------------------------------------- kappa
def kappa_cohen(a: pd.Series, b: pd.Series) -> float:
    """Concordancia entre dos anotadores descontando el azar."""
    par = pd.DataFrame({"a": a, "b": b}).replace("", pd.NA).dropna()
    if par.empty:
        return float("nan")
    n = len(par)
    observado = (par["a"] == par["b"]).mean()
    categorias = sorted(set(par["a"]) | set(par["b"]))
    esperado = sum((par["a"] == c).mean() * (par["b"] == c).mean() for c in categorias)
    if esperado >= 1.0:
        return float("nan")
    return (observado - esperado) / (1 - esperado)


def kappa_fleiss(df: pd.DataFrame) -> float:
    """Concordancia entre tres o mas anotadores sobre los items completos."""
    datos = df.replace("", pd.NA).dropna()
    if datos.empty or datos.shape[1] < 2:
        return float("nan")
    categorias = sorted({v for col in datos.columns for v in datos[col].unique()})
    n_anot = datos.shape[1]
    conteos = pd.DataFrame(
        {c: (datos == c).sum(axis=1) for c in categorias}, index=datos.index)

    p_i = ((conteos ** 2).sum(axis=1) - n_anot) / (n_anot * (n_anot - 1))
    p_barra = p_i.mean()
    p_j = conteos.sum(axis=0) / (len(datos) * n_anot)
    p_e = (p_j ** 2).sum()
    if p_e >= 1.0:
        return float("nan")
    return (p_barra - p_e) / (1 - p_e)


def interpretar(k: float) -> str:
    if pd.isna(k):
        return "no calculable"
    if k < 0.20:
        return "muy baja"
    if k < 0.40:
        return "baja"
    if k < 0.60:
        return "moderada"
    if k < 0.80:
        return "sustancial"
    return "casi perfecta"


# ---------------------------------------------------------------- consenso
def consenso(df: pd.DataFrame) -> pd.Series:
    """
    Estandar de oro por mayoria simple.

    Un item queda sin resolver cuando no hay mayoria o cuando la mayoria es
    DUDOSO. Esos casos se reportan aparte en vez de forzar una etiqueta.
    """
    resultado = {}
    for ident, fila in df.iterrows():
        votos = [v for v in fila if v in ("SI", "NO", "DUDOSO")]
        if not votos:
            resultado[ident] = ""
            continue
        conteo = pd.Series(votos).value_counts()
        if conteo.iloc[0] * 2 > len(votos) and conteo.index[0] in ("SI", "NO"):
            resultado[ident] = conteo.index[0]
        else:
            resultado[ident] = "SIN_RESOLVER"
    return pd.Series(resultado)


# ------------------------------------------------------------ conciliacion
# Decisiones de la sesion de conciliacion, por hoja. Vacio si no se usa el
# modo --conciliado.
CONCILIADO: dict = {}
LIBRO_CONCILIACION = CARPETA / "Conciliacion_Anotadores.xlsx"


def cargar_conciliacion() -> dict:
    """
    Lee las decisiones del libro de conciliacion.

    Solo alimentan el estandar de oro. El kappa se sigue calculando sobre las
    etiquetas que cada anotador emitio por separado: recalcularlo sobre
    decisiones tomadas en grupo mostraria acuerdo perfecto en los items
    discutidos y subiria el valor por construccion, sin que la concordancia
    real haya cambiado.
    """
    wb = load_workbook(LIBRO_CONCILIACION, data_only=True)
    out = {}
    for hoja in ("A_Entidades", "B_Intencion"):
        h = wb[hoja]
        decisiones, invalidas = {}, []
        for fila in range(6, h.max_row + 1):
            ident = h.cell(row=fila, column=1).value
            if ident is None:
                continue
            valor = h.cell(row=fila, column=7).value
            etq = str(valor).strip().upper() if valor is not None else ""
            if etq in ("SI", "NO"):
                decisiones[str(ident).strip()] = etq
            else:
                invalidas.append(str(ident))
        if invalidas:
            raise ValueError(f"{hoja}: {len(invalidas)} items sin decision SI/NO "
                             f"({', '.join(invalidas[:5])}...). La conciliacion "
                             f"debe resolverlos todos antes de evaluar.")
        out[hoja] = decisiones
    return out


def aplicar_conciliacion(oro: pd.Series, hoja: str) -> pd.Series:
    if not CONCILIADO:
        return oro
    oro = oro.copy()
    for ident, etq in CONCILIADO[hoja].items():
        if ident in oro.index:
            oro[ident] = etq
    return oro


# ---------------------------------------------------------------- informe
def bloque_concordancia(df: pd.DataFrame, titulo: str, lineas: list):
    lineas.append(f"\n### Concordancia — {titulo}")
    completos = df.replace("", pd.NA).dropna()
    lineas.append(f"- Items con los tres anotadores: {len(completos)} de {len(df)}")
    for a, b in itertools.combinations(df.columns, 2):
        k = kappa_cohen(df[a], df[b])
        lineas.append(f"- Cohen {a.split('_')[0]} vs {b.split('_')[0]}: "
                      f"{k:.3f} ({interpretar(k)})")
    kf = kappa_fleiss(df)
    lineas.append(f"- **Fleiss (los tres): {kf:.3f} ({interpretar(kf)})**")
    if not pd.isna(kf) and kf < 0.70:
        lineas.append("- Por debajo de la meta de 0,70: conviene revisar el manual "
                      "de anotacion y conciliar los criterios antes de reportar.")
    return kf


def evaluar_entidades(lineas: list):
    lineas.append("\n## Tarea A — Precision del reconocimiento de entidades")
    df = cargar("A_Entidades")
    if df.empty:
        lineas.append("\nNo hay anotaciones todavia.")
        return
    anotados = (df != "").sum()
    lineas.append(f"\nAvance por anotador: "
                  + ", ".join(f"{c.split('_')[0]} {n}/{len(df)}"
                              for c, n in anotados.items()))
    bloque_concordancia(df, "entidades", lineas)

    oro = aplicar_conciliacion(consenso(df), "A_Entidades")
    resueltos = oro[oro.isin(["SI", "NO"])]
    if resueltos.empty:
        lineas.append("\nSin items resueltos por mayoria.")
        return
    positivos = int((resueltos == "SI").sum())
    precision = positivos / len(resueltos)
    lineas.append(f"\n- Entidades resueltas por mayoria: {len(resueltos)}")
    lineas.append(f"- Sin resolver o dudosas: {int((oro == 'SIN_RESOLVER').sum())}")
    lineas.append(f"- Entidades que SI son producto: {positivos}")
    lineas.append(f"- **Precision del extractor: {precision:.1%}**")
    lineas.append("\nEs precision, no recall: mide que proporcion de lo que el "
                  "sistema propuso era correcto. El recall se estima con la tarea C.")


def evaluar_intencion(lineas: list):
    lineas.append("\n## Tarea B — Desempeño del clasificador de intencion")
    df = cargar("B_Intencion")
    if df.empty:
        lineas.append("\nNo hay anotaciones todavia.")
        return

    clave_path = CARPETA / "clave_maestra_comentarios.csv"
    estratos_path = CARPETA / "estratos.csv"
    if not clave_path.exists() or not estratos_path.exists():
        lineas.append("\nFaltan las claves maestras; no se puede evaluar.")
        return
    clave = pd.read_csv(clave_path).set_index("id")
    estratos = pd.read_csv(estratos_path).set_index("estrato")

    anotados = (df != "").sum()
    lineas.append(f"\nAvance por anotador: "
                  + ", ".join(f"{c.split('_')[0]} {n}/{len(df)}"
                              for c, n in anotados.items()))
    bloque_concordancia(df, "intencion de compra", lineas)

    oro = aplicar_conciliacion(consenso(df), "B_Intencion")
    tabla = pd.DataFrame({"oro": oro}).join(clave[["prediccion_sistema"]], how="inner")
    tabla = tabla[tabla["oro"].isin(["SI", "NO"])]
    if tabla.empty:
        lineas.append("\nSin items resueltos por mayoria.")
        return

    marcados = tabla[tabla["prediccion_sistema"] == "SI"]
    no_marcados = tabla[tabla["prediccion_sistema"] == "NO"]
    vp = int((marcados["oro"] == "SI").sum())
    fp = int((marcados["oro"] == "NO").sum())

    lineas.append(f"\n### Estrato marcado por el sistema ({len(marcados)} anotados)")
    lineas.append(f"- Verdaderos positivos: {vp}")
    lineas.append(f"- Falsos positivos: {fp}")
    precision = vp / (vp + fp) if (vp + fp) else float("nan")
    lineas.append(f"- **Precision: {precision:.1%}**")

    if no_marcados.empty:
        lineas.append("\nSin datos del estrato no marcado: no se puede estimar recall.")
        return

    # El estrato no marcado se muestreo, asi que hay que reponderar a poblacion.
    pob_no = int(estratos.loc["no_marcados", "poblacion"])
    mues_no = len(no_marcados)
    tasa_fn = (no_marcados["oro"] == "SI").mean()
    fn_estimados = tasa_fn * pob_no

    lineas.append(f"\n### Estrato no marcado ({mues_no} anotados de {pob_no})")
    lineas.append(f"- Con intencion segun los anotadores: "
                  f"{int((no_marcados['oro'] == 'SI').sum())} ({tasa_fn:.1%})")
    lineas.append(f"- Falsos negativos estimados en la poblacion: {fn_estimados:.0f}")

    recall = vp / (vp + fn_estimados) if (vp + fn_estimados) else float("nan")
    f1 = (2 * precision * recall / (precision + recall)
          if precision and recall and (precision + recall) else float("nan"))

    lineas.append("\n### Metricas del objetivo especifico 2")
    lineas.append(f"- Precision : {precision:.3f}")
    lineas.append(f"- Recall    : {recall:.3f}  (estimado por reponderacion)")
    lineas.append(f"- **F1-score: {f1:.3f}**")
    meta = "CUMPLE" if f1 >= 0.80 else "NO CUMPLE"
    lineas.append(f"- Meta comprometida F1 >= 0,80: **{meta}**")

    lineas.append("\n### Matriz de confusion (reponderada a poblacion)")
    vn_est = (1 - tasa_fn) * pob_no
    lineas.append("")
    lineas.append("| | Sistema dice SI | Sistema dice NO |")
    lineas.append("|---|---|---|")
    lineas.append(f"| Anotadores dicen SI | {vp} | {fn_estimados:.0f} |")
    lineas.append(f"| Anotadores dicen NO | {fp} | {vn_est:.0f} |")


def main() -> int:
    parser = argparse.ArgumentParser(description="Evalua la anotacion manual")
    parser.add_argument("--salida", default=None)
    parser.add_argument("--conciliado", action="store_true",
                        help="Usa las decisiones de Conciliacion_Anotadores.xlsx "
                             "como estandar de oro. El kappa no cambia.")
    args = parser.parse_args()
    if args.salida is None:
        args.salida = ("anotacion/resultados_conciliados.md" if args.conciliado
                       else "anotacion/resultados_anotacion.md")

    if not CARPETA.exists():
        print("No existe la carpeta anotacion/. Ejecute construir_anotacion.py.")
        return 1

    if args.conciliado:
        if not LIBRO_CONCILIACION.exists():
            print(f"No existe {LIBRO_CONCILIACION}.")
            return 1
        CONCILIADO.update(cargar_conciliacion())

    lineas = ["# Resultados de la anotacion manual"
              + (" — estandar de oro conciliado" if args.conciliado else ""),
              "",
              "Generado por evaluar_anotacion.py. Las metricas de esta seccion "
              "alimentan el numeral 7.4 y el capitulo 8 del documento."]
    if args.conciliado:
        lineas.append("\nEl estandar de oro incorpora las decisiones de la sesion "
                      "de conciliacion. El kappa se calcula sobre las etiquetas "
                      "independientes originales y por eso no cambia.")
    evaluar_entidades(lineas)
    evaluar_intencion(lineas)

    texto = "\n".join(lineas)
    print(texto)
    Path(args.salida).write_text(texto, encoding="utf-8")
    print(f"\n\nInforme guardado en {args.salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
