"""
Experimento de convergencia entre señal social y demanda de busqueda (OE3).

Responde la pregunta del objetivo especifico 3 y contrasta la hipotesis H1:
¿la velocidad de interaccion en redes sociales se asocia con el crecimiento del
volumen de busqueda, y con cuanto rezago llega ese interes a Colombia?

Diferencia importante frente al pipeline: aqui NO se aplica el umbral de
validacion. El umbral existe para decidir que productos se recomiendan; para
estudiar la correlacion hacen falta tambien los que no crecen, porque filtrar
antes de correlacionar seleccionaria la muestra sobre la propia variable
dependiente y sesgaria el resultado.

Uso:
    ./venv/Scripts/python.exe experimento_correlacion.py
    ./venv/Scripts/python.exe experimento_correlacion.py --todas   # sin anotacion
    ./venv/Scripts/python.exe experimento_correlacion.py --limite 30
"""
import argparse
import json
import sqlite3
import sys
from pathlib import Path

import pandas as pd
from scipy import stats

from src.cross_validation import TrendValidator
from src.database import StorageAndScoring

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CARPETA = Path("anotacion")
CACHE = CARPETA / "cache_trends.csv"
SALIDA_TABLA = CARPETA / "datos_correlacion.csv"
SALIDA_INFORME = CARPETA / "resultados_correlacion.md"
ANOTADORES = ["Daniela_Trujillo", "Miguel_Reina", "Gabriel_Zambrano"]


# ------------------------------------------------------- productos confirmados
def productos_confirmados() -> list:
    """Entidades que la mayoria de los anotadores marco como producto real."""
    from openpyxl import load_workbook

    votos = {}
    libros = 0
    for nombre in ANOTADORES:
        ruta = CARPETA / f"Anotacion_{nombre}.xlsx"
        if not ruta.exists():
            continue
        libros += 1
        hoja = load_workbook(ruta, data_only=True)["A_Entidades"]
        for fila in range(7, hoja.max_row + 1):
            entidad = hoja.cell(row=fila, column=2).value
            etiqueta = hoja.cell(row=fila, column=3).value
            if not entidad:
                continue
            etiqueta = str(etiqueta).strip().upper() if etiqueta else ""
            votos.setdefault(entidad, []).append(etiqueta)

    if libros == 0:
        return []

    confirmados = []
    for entidad, marcas in votos.items():
        validas = [m for m in marcas if m in ("SI", "NO", "DUDOSO")]
        if validas and validas.count("SI") * 2 > len(validas):
            confirmados.append(entidad)
    return sorted(confirmados)


def todas_las_entidades() -> list:
    con = sqlite3.connect("data/trends.db")
    ents = set()
    for (j,) in con.execute("SELECT extracted_products FROM processed_entities"):
        ents.update(json.loads(j or "[]"))
    con.close()
    return sorted(ents)


# ------------------------------------------------------- metricas sociales
def metricas_sociales(productos: list) -> pd.DataFrame:
    """Velocidad de interaccion, menciones e intencion por producto."""
    con = sqlite3.connect("data/trends.db")
    df = pd.read_sql(
        "SELECT p.extracted_products, p.purchase_intent_score, "
        "       r.likes_count, r.shares_count, r.timestamp "
        "FROM processed_entities p JOIN raw_social_data r "
        "  ON p.video_id = r.video_id", con)
    con.close()

    ahora = pd.Timestamp.now(tz="UTC")
    objetivo = {p.lower() for p in productos}
    acumulado = {}
    for _, fila in df.iterrows():
        velocidad = StorageAndScoring._engagement_velocity(fila, ahora)
        intent = float(fila.get("purchase_intent_score") or 0.0)
        for p in json.loads(fila["extracted_products"] or "[]"):
            clave = p.lower()
            if clave not in objetivo:
                continue
            reg = acumulado.setdefault(clave, {"velocidad": 0.0, "intencion": [],
                                               "menciones": 0})
            reg["velocidad"] += velocidad
            reg["intencion"].append(intent)
            reg["menciones"] += 1

    filas = [{"producto": k,
              "velocidad_social": v["velocidad"],
              "menciones": v["menciones"],
              "intencion_media": sum(v["intencion"]) / len(v["intencion"])}
             for k, v in acumulado.items()]
    return pd.DataFrame(filas)


# ------------------------------------------------------- consulta a Trends
def consultar_trends(productos: list, validator: TrendValidator) -> pd.DataFrame:
    """
    Consulta Google Trends con cache en disco.

    Cada termino tarda unos 15 segundos por geografia. El cache permite reanudar
    si Google corta la conexion a mitad de camino, sin repetir lo ya obtenido.
    """
    cache = {}
    if CACHE.exists():
        previo = pd.read_csv(CACHE)
        cache = {r["producto"]: r.to_dict() for _, r in previo.iterrows()}
        print(f"Cache: {len(cache)} terminos ya consultados.")

    resultados = list(cache.values())
    pendientes = [p for p in productos if p not in cache]
    print(f"Terminos por consultar: {len(pendientes)}")

    for i, producto in enumerate(pendientes, start=1):
        print(f"  [{i}/{len(pendientes)}] {producto}")
        registro = {"producto": producto, "con_datos": 0,
                    "crecimiento_mundial": None, "volumen_mundial": None,
                    "crecimiento_local": None, "volumen_local": None,
                    "rezago_dias": None, "correlacion_rezago": None}

        df_mundo = validator.fetch_interest_over_time(producto, geo=validator.geo)
        if df_mundo is not None and producto in df_mundo.columns:
            _, crec, vol = validator.calculate_growth_metrics(df_mundo, producto)
            registro.update(con_datos=1, crecimiento_mundial=crec,
                            volumen_mundial=vol)

            if validator.compare_geos:
                df_local = validator.fetch_interest_over_time(
                    producto, geo=validator.geo_local)
                if df_local is not None and producto in df_local.columns:
                    if df_mundo[producto].equals(df_local[producto]):
                        # Si ambas geografias devuelven la misma serie, el filtro
                        # geografico no se aplico y comparar no significa nada.
                        print("      aviso: series identicas en ambas geografias; "
                              "se descarta la comparacion local")
                    else:
                        _, crec_l, vol_l = validator.calculate_growth_metrics(
                            df_local, producto)
                        registro.update(crecimiento_local=crec_l,
                                        volumen_local=vol_l)
                        rezago, corr = validator.estimate_lag_days(
                            df_mundo[producto], df_local[producto])
                        registro.update(rezago_dias=rezago,
                                        correlacion_rezago=corr)

        resultados.append(registro)
        # Se guarda en cada iteracion: una caida no cuesta lo ya consultado.
        pd.DataFrame(resultados).to_csv(CACHE, index=False, encoding="utf-8-sig")

    return pd.DataFrame(resultados)


# ------------------------------------------------------- analisis
def spearman(x: pd.Series, y: pd.Series):
    par = pd.DataFrame({"x": x, "y": y}).dropna()
    if len(par) < 5 or par["x"].nunique() < 3 or par["y"].nunique() < 3:
        return None, None, len(par)
    rho, p = stats.spearmanr(par["x"], par["y"])
    return rho, p, len(par)


def linea_correlacion(nombre, x, y, lineas):
    rho, p, n = spearman(x, y)
    if rho is None:
        lineas.append(f"| {nombre} | n insuficiente ({n}) | — | — |")
        return
    signif = "si" if p < 0.05 else "no"
    lineas.append(f"| {nombre} | {rho:+.3f} | {p:.4f} | {n} | {signif} |")


def main() -> int:
    parser = argparse.ArgumentParser(description="Experimento de correlacion OE3")
    parser.add_argument("--todas", action="store_true",
                        help="Usa todas las entidades, sin esperar la anotacion")
    parser.add_argument("--limite", type=int, default=60,
                        help="Maximo de terminos a consultar")
    parser.add_argument("--geo", default=None)
    parser.add_argument("--geo-local", dest="geo_local", default=None)
    args = parser.parse_args()

    CARPETA.mkdir(exist_ok=True)

    if args.todas:
        productos = todas_las_entidades()
        origen = "todas las entidades extraidas (sin filtrar por anotacion)"
    else:
        productos = productos_confirmados()
        origen = "entidades confirmadas como producto por mayoria de anotadores"
        if not productos:
            print("No hay anotaciones completas todavia.")
            print("Complete los libros o ejecute con --todas para una corrida "
                  "exploratoria.")
            return 1

    productos = productos[:args.limite]
    print(f"Productos a analizar: {len(productos)}  ({origen})\n")

    social = metricas_sociales(productos)
    if social.empty:
        print("No se pudieron calcular metricas sociales.")
        return 1

    validator = TrendValidator(geo=args.geo, geo_local=args.geo_local)
    trends = consultar_trends(productos, validator)

    datos = social.merge(trends, on="producto", how="inner")
    datos.to_csv(SALIDA_TABLA, index=False, encoding="utf-8-sig")

    con_datos = datos[datos["con_datos"] == 1]

    lineas = ["# Convergencia entre señal social y demanda de busqueda",
              "",
              "Resultados del objetivo especifico 3 y contraste de la hipotesis H1.",
              f"Generado por experimento_correlacion.py sobre {origen}.",
              "",
              "## Cobertura",
              f"- Productos analizados: {len(datos)}",
              f"- Con datos en Google Trends: {len(con_datos)} "
              f"({100*len(con_datos)/max(len(datos),1):.0f}%)",
              f"- Sin datos: {len(datos) - len(con_datos)}",
              "",
              "La proporcion sin datos es en si misma un resultado: Google Trends "
              "solo cubre terminos con volumen de busqueda apreciable, de modo que "
              "los productos mas incipientes son justamente los que no puede "
              "validar. Es una limitacion estructural del metodo de validacion.",
              "",
              "## Correlaciones de Spearman",
              "",
              "| Relacion | rho | valor p | n | significativa |",
              "|---|---|---|---|---|"]

    linea_correlacion("Velocidad social vs crecimiento de busqueda mundial",
                      con_datos["velocidad_social"],
                      con_datos["crecimiento_mundial"], lineas)
    linea_correlacion("Menciones vs volumen de busqueda mundial",
                      con_datos["menciones"], con_datos["volumen_mundial"], lineas)
    linea_correlacion("Intencion de compra vs crecimiento de busqueda",
                      con_datos["intencion_media"],
                      con_datos["crecimiento_mundial"], lineas)
    linea_correlacion("Velocidad social vs crecimiento de busqueda en Colombia",
                      con_datos["velocidad_social"],
                      con_datos["crecimiento_local"], lineas)

    lineas += ["",
               "El contraste de H1 se decide sobre la primera fila: la hipotesis "
               "postula una asociacion monotonica positiva con rho de al menos "
               "0,40 y p inferior a 0,05."]

    con_rezago = con_datos[con_datos["rezago_dias"].notna()]
    lineas += ["", "## Rezago de adopcion de Colombia frente al mercado mundial", ""]
    if con_rezago.empty:
        lineas.append("Ningun producto alcanzo la correlacion minima exigida para "
                      "reportar un rezago. Con series planas o muy cortas la "
                      "propagacion no es distinguible del ruido, y el metodo "
                      "prefiere no reportar antes que inventar un valor.")
    else:
        lineas += [
            f"- Productos con rezago medible: {len(con_rezago)} de {len(con_datos)}",
            f"- Rezago mediano: {con_rezago['rezago_dias'].median():.0f} dias",
            f"- Rango: {con_rezago['rezago_dias'].min():.0f} a "
            f"{con_rezago['rezago_dias'].max():.0f} dias",
            f"- Correlacion mediana en el rezago: "
            f"{con_rezago['correlacion_rezago'].median():.3f}",
            "",
            "Ese rezago mediano es la ventana operativa: los dias de margen que "
            "tiene una PYME colombiana para abastecerse antes de que la demanda "
            "local se manifieste.",
        ]

    brecha = con_datos.dropna(subset=["crecimiento_mundial", "crecimiento_local"])
    if not brecha.empty:
        b = brecha["crecimiento_mundial"] - brecha["crecimiento_local"]
        lineas += ["", "## Brecha de adopcion", "",
                   f"- Productos comparables: {len(brecha)}",
                   f"- Brecha mediana (mundial menos local): {b.median():+.1%}",
                   f"- Con brecha positiva (aun no llega a Colombia): "
                   f"{int((b > 0).sum())} de {len(brecha)}"]

    lineas += ["", "## Tabla de datos", "",
               f"El detalle por producto quedo en `{SALIDA_TABLA}`, listo para "
               "anexarse al documento."]

    texto = "\n".join(lineas)
    print("\n" + texto)
    SALIDA_INFORME.write_text(texto, encoding="utf-8")
    print(f"\n\nInforme guardado en {SALIDA_INFORME}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
