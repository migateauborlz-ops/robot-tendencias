# -*- coding: utf-8 -*-
"""
Construye el libro de trabajo para la sesion de conciliacion.

Reune los items en que los tres anotadores no coincidieron de forma unanime, o
en que alguno marco DUDOSO, y los presenta con la seleccion de cada uno para
resolverlos en grupo.

Dos decisiones de diseno que conviene no cambiar:

1. No se muestra la prediccion del sistema. El muestreo original la oculto a
   proposito para que el anotador no pudiera inferir a que estrato pertenecia
   cada fila. Mostrarla ahora sesgaria el estandar de oro hacia el sistema y
   volveria optimista toda metrica calculada contra el.

2. No se sugiere una respuesta. El libro agrupa y ordena, pero no propone
   etiqueta. El valor de la conciliacion esta en que la decision salga del
   protocolo discutido y no de una sugerencia automatica.

Uso:
    ./venv/Scripts/python.exe construir_conciliacion.py
"""
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

CARPETA = Path("anotacion")
SALIDA = CARPETA / "Conciliacion_Anotadores.xlsx"
ANOTADORES = [("Daniela", "Anotacion_Daniela_Trujillo.xlsx"),
              ("Miguel", "Anotacion_Miguel_Reina.xlsx"),
              ("Gabriel", "Anotacion_Gabriel_Zambrano.xlsx")]
PRIMERA_FILA = 7
VALIDAS = {"SI", "NO", "DUDOSO"}

GRANATE = "8C1D1D"
CREMA = "F4F1EA"
AMARILLO = "FFF3C4"
GRIS = "E8E8E8"
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def leer(ruta, hoja):
    wb = load_workbook(ruta, data_only=True)
    h = wb[hoja]
    out = {}
    for f in range(PRIMERA_FILA, h.max_row + 1):
        ident = h.cell(row=f, column=1).value
        if ident is None:
            continue
        texto = h.cell(row=f, column=2).value or ""
        valor = h.cell(row=f, column=3).value
        etq = str(valor).strip().upper() if valor is not None else ""
        out[str(ident).strip()] = (etq if etq in VALIDAS else "",
                                   str(texto).replace("\n", " ").strip())
    return out


def clasificar(vs):
    """Tipo de desacuerdo a partir de las tres etiquetas."""
    c = Counter(vs)
    dudas = c["DUDOSO"]
    if dudas == 3:
        return "Duda unanime"
    if dudas and c["SI"] and c["NO"]:
        return "Duda y division"
    if dudas:
        return "Duda parcial"
    return "Division 2-1"


def recolectar(hoja):
    datos = {n: leer(CARPETA / f, hoja) for n, f in ANOTADORES}
    ids = sorted(datos["Daniela"])
    filas = []
    for i in ids:
        vs = [datos[n].get(i, ("", ""))[0] for n, _ in ANOTADORES]
        texto = datos["Daniela"].get(i, ("", ""))[1]
        unanime = len(set(vs)) == 1 and vs[0] != "DUDOSO"
        if unanime:
            continue
        filas.append({"id": i, "texto": texto, "vs": vs,
                      "tipo": clasificar(vs)})
    # Se agrupa por tipo y, dentro de cada grupo, por texto. El orden
    # alfabetico junta los casos parecidos, que suelen resolverse con la
    # misma regla y de una sola vez.
    orden = {"Duda unanime": 0, "Duda y division": 1,
             "Duda parcial": 2, "Division 2-1": 3}
    filas.sort(key=lambda r: (orden[r["tipo"]], r["texto"].lower()))
    return filas, len(ids)


def encabezado(ws, titulo, subtitulo, ancho):
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=14, color=GRANATE)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, size=10)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=ancho)
    ws.row_dimensions[2].height = 30


def hoja_items(wb, nombre, titulo, pregunta, filas, total):
    ws = wb.create_sheet(nombre)
    cols = ["ID", "Texto", "Daniela", "Miguel", "Gabriel", "Tipo de desacuerdo",
            "DECISION", "Regla que la sustenta"]
    anchos = [8, 62, 11, 11, 11, 20, 13, 40]
    encabezado(ws, titulo,
               "%d items por conciliar de %d anotados. %s  "
               "Complete unicamente las columnas DECISION y Regla."
               % (len(filas), total, pregunta), len(cols))

    fila0 = 5
    for c, (nom, an) in enumerate(zip(cols, anchos), 1):
        celda = ws.cell(row=fila0, column=c, value=nom)
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=GRANATE)
        celda.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = an

    dv = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True,
                        showErrorMessage=True,
                        error="Escriba SI o NO. La conciliacion debe resolver "
                              "la duda, no conservarla.")
    ws.add_data_validation(dv)

    tipo_previo = None
    for n, r in enumerate(filas):
        f = fila0 + 1 + n
        ws.cell(row=f, column=1, value=r["id"])
        ws.cell(row=f, column=2, value=r["texto"][:300])
        for k, v in enumerate(r["vs"]):
            celda = ws.cell(row=f, column=3 + k, value=v)
            celda.alignment = Alignment(horizontal="center")
            if v == "DUDOSO":
                celda.fill = PatternFill("solid", fgColor=AMARILLO)
                celda.font = Font(bold=True)
        ws.cell(row=f, column=6, value=r["tipo"])
        decision = ws.cell(row=f, column=7)
        decision.fill = PatternFill("solid", fgColor=AMARILLO)
        decision.alignment = Alignment(horizontal="center")
        dv.add(decision)
        for c in range(1, len(cols) + 1):
            ws.cell(row=f, column=c).border = BORDE
            if c == 2:
                ws.cell(row=f, column=c).alignment = Alignment(wrap_text=True,
                                                               vertical="top")
        # Una linea mas marcada separa los bloques por tipo de desacuerdo.
        if tipo_previo is not None and r["tipo"] != tipo_previo:
            for c in range(1, len(cols) + 1):
                ws.cell(row=f, column=c).border = Border(
                    top=Side(style="medium", color=GRANATE),
                    left=Side(style="thin", color="BFBFBF"),
                    right=Side(style="thin", color="BFBFBF"),
                    bottom=Side(style="thin", color="BFBFBF"))
        tipo_previo = r["tipo"]

    ws.freeze_panes = ws.cell(row=fila0 + 1, column=3)
    return len(filas)


def hoja_instrucciones(wb, n_ent, n_int, tot_ent, tot_int):
    ws = wb.create_sheet("Instrucciones", 0)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 108
    lineas = [
        ("t", "Sesion de conciliacion de la anotacion"),
        ("s", "%d items de entidades y %d de intencion, sobre %d y %d anotados."
              % (n_ent, n_int, tot_ent, tot_int)),
        ("", ""),
        ("h", "Que resuelve esta sesion y que no"),
        ("p", "Resuelve el estandar de oro. Las etiquetas conciliadas son las "
              "que se usan para calcular precision, recall y F1-score del "
              "sistema. Ese calculo si queda valido despues de la sesion."),
        ("p", "NO produce un kappa nuevo valido. El kappa mide acuerdo entre "
              "anotadores que trabajaron por separado. Si se recalcula sobre "
              "las etiquetas conciliadas, los items discutidos apareceran con "
              "acuerdo perfecto y el valor subira por construccion, no por "
              "mejora real. Reportar ese numero como concordancia "
              "interanotador seria incorrecto y un jurado puede senalarlo."),
        ("p", "El kappa medido (0,519 en entidades y 0,612 en intencion) se "
              "reporta tal como se obtuvo. Para exhibir un kappa que supere el "
              "umbral de 0,70 hace falta corregir el protocolo con lo que "
              "salga de esta sesion y volver a anotar una muestra nueva de "
              "forma independiente."),
        ("", ""),
        ("h", "Los dos problemas que hay que resolver"),
        ("p", "1. Cuando corresponde marcar DUDOSO. El uso fue desigual: un "
              "anotador lo empleo 59 veces en intencion y otro 19, sobre el "
              "mismo material. Hay que acordar un criterio y escribirlo."),
        ("p", "2. Donde esta la frontera entre categoria generica y producto "
              "concreto. Los desacuerdos reales en entidades se concentran en "
              "casos como categorias de actividad frente a objetos."),
        ("", ""),
        ("h", "Como trabajar"),
        ("p", "Los items vienen agrupados por tipo de desacuerdo y ordenados "
              "alfabeticamente dentro de cada grupo. Los casos parecidos "
              "quedan juntos a proposito, porque suelen resolverse con la "
              "misma regla y de una sola vez."),
        ("p", "Empiecen por el bloque Duda unanime. Son los casos donde el "
              "protocolo fallo para los tres, y de ahi salen las reglas mas "
              "utiles."),
        ("p", "Cada decision debe quedar en SI o en NO. No se permite dejar "
              "DUDOSO: el proposito de la sesion es resolverlo."),
        ("p", "Anoten en la ultima columna que regla sustenta la decision. Si "
              "la regla no existe todavia en el manual, registrenla en la hoja "
              "Reglas nuevas. Esa hoja es el producto mas valioso de la "
              "sesion, porque es lo que permite una segunda ronda con kappa "
              "valido."),
        ("", ""),
        ("h", "Lo que este libro no muestra, y por que"),
        ("p", "No aparece la prediccion del sistema. El muestreo original la "
              "oculto para que el anotador no supiera a que estrato pertenecia "
              "cada fila. Mostrarla ahora inclinaria el estandar de oro hacia "
              "el sistema y volveria optimista cualquier metrica calculada "
              "contra el."),
        ("p", "Tampoco hay etiqueta sugerida. La decision tiene que salir del "
              "protocolo discutido entre ustedes."),
        ("", ""),
        ("h", "Al terminar"),
        ("p", "Guarden el archivo y ejecuten: "
              "./venv/Scripts/python.exe evaluar_anotacion.py --conciliado"),
    ]
    f = 2
    for tipo, texto in lineas:
        c = ws.cell(row=f, column=2, value=texto)
        if tipo == "t":
            c.font = Font(bold=True, size=15, color=GRANATE)
        elif tipo == "s":
            c.font = Font(italic=True, size=10)
        elif tipo == "h":
            c.font = Font(bold=True, size=11, color=GRANATE)
        else:
            c.font = Font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[f].height = max(14, 13 * (len(texto) // 100 + 1))
        f += 1
    return ws


def hoja_reglas(wb):
    ws = wb.create_sheet("Reglas nuevas")
    encabezado(ws, "Reglas acordadas en la conciliacion",
               "Cada regla que resuelva un grupo de casos se registra aqui. "
               "Esta hoja se traslada al manual de anotacion y habilita una "
               "segunda ronda con kappa valido.", 4)
    cols = ["No.", "Tarea", "Regla acordada", "Ejemplo del corpus"]
    anchos = [6, 14, 64, 42]
    for c, (nom, an) in enumerate(zip(cols, anchos), 1):
        celda = ws.cell(row=5, column=c, value=nom)
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=GRANATE)
        celda.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = an
    for f in range(6, 31):
        ws.cell(row=f, column=1, value=f - 5).alignment = Alignment(
            horizontal="center")
        for c in range(1, 5):
            ws.cell(row=f, column=c).border = BORDE
            if c >= 3:
                ws.cell(row=f, column=c).fill = PatternFill("solid",
                                                            fgColor=AMARILLO)
    ws.freeze_panes = "A6"


def main():
    if not CARPETA.exists():
        print("No existe la carpeta anotacion/.")
        return 1

    ent, tot_ent = recolectar("A_Entidades")
    inte, tot_int = recolectar("B_Intencion")

    wb = Workbook()
    wb.remove(wb.active)
    hoja_instrucciones(wb, len(ent), len(inte), tot_ent, tot_int)
    hoja_items(wb, "A_Entidades", "Tarea A — Entidades por conciliar",
               "Pregunta: ¿el texto nombra un producto fisico que una PYME "
               "podria comprar y revender?", ent, tot_ent)
    hoja_items(wb, "B_Intencion", "Tarea B — Intencion por conciliar",
               "Pregunta: ¿el comentario expresa intencion de compra?",
               inte, tot_int)
    hoja_reglas(wb)
    wb.save(SALIDA)

    print("=" * 70)
    print("Libro de conciliacion generado")
    print("=" * 70)
    for etq, filas, total in (("Entidades", ent, tot_ent),
                              ("Intencion", inte, tot_int)):
        c = Counter(r["tipo"] for r in filas)
        print("\n%s: %d por conciliar de %d (%.0f%%)"
              % (etq, len(filas), total, len(filas) / total * 100))
        for tipo in ("Duda unanime", "Duda y division", "Duda parcial",
                     "Division 2-1"):
            if c[tipo]:
                print("   %-18s %3d" % (tipo, c[tipo]))
    print("\narchivo: %s" % SALIDA)
    return 0


if __name__ == "__main__":
    sys.exit(main())
